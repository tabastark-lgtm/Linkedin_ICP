import copy
import hashlib
import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError, URLError
from openpyxl import load_workbook
from pesquisa.core import domain, company_url, validate_record, empty_record
from pesquisa.excel import inspect_book, write_output
from pesquisa.store import Store
from pesquisa.native import handle, read_message, write_message
from pesquisa.search import search, SearchError

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.store = Store(self.root / "dados")
        self.source = EXAMPLES / "01_empresas_ficticias.xlsx"
        self.output = self.root / "saida.xlsx"
        self.job = self.store.import_job(self.source, "Empresas", self.output)
        self.record = self.store.records(self.job)[0]

    def tearDown(self): self.temp.cleanup()

    def test_website_validation(self):
        self.assertEqual(domain("HTTPS://www.Exemplo.com.br/sobre?q=1"), "exemplo.com.br")
        self.assertEqual(domain("sub.exemplo.com"), "sub.exemplo.com")
        for value in ("", "email@exemplo.com", "não é URL", "javascript:alert(1)", "https://a.com@b.com", "localhost", "http://a.com:80"):
            with self.subTest(value=value), self.assertRaises(ValueError): domain(value)

    def test_linkedin_validation(self):
        self.assertEqual(company_url("https://br.linkedin.com/company/teste/about/?x=1"), "https://www.linkedin.com/company/teste/")
        for value in ("https://linkedin.com.evil.example/company/test/", "https://evil-linkedin.com/company/test", "https://linkedin.com/in/test", "https://linkedin.com/company/", "https://x@linkedin.com/company/test", "http://linkedin.com/company/test"):
            with self.subTest(value=value), self.assertRaises(ValueError): company_url(value)

    def test_import_limits_and_headers(self):
        self.assertEqual(len(inspect_book(EXAMPLES / "04_limite_50.xlsx", "Empresas")["records"]), 50)
        for filename in ("02_sem_cabecalho.xlsx", "03_cabecalho_duplicado.xlsx", "04_limite_51.xlsx", "05_formula_sem_resultado.xlsx"):
            with self.subTest(filename=filename), self.assertRaises(ValueError): inspect_book(EXAMPLES / filename, "Empresas")
        self.assertEqual(inspect_book(EXAMPLES / "06_formula_com_resultado.xlsx", "Empresas")["rows"][1][1], 2)

    def test_export_preserves_order_duplicates_and_original(self):
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        d = self.record["data"]; d.update(name="Aurora TESTE", members="123", size="51–200 funcionários", notes="=1+1", url="https://www.linkedin.com/company/teste/", confirmed=True, consulted_at="2026-09-29T15:00:00-03:00")
        self.store.save(self.record["id"], d)
        write_output(self.store.job(self.job), self.store.records(self.job), self.output)
        wb = load_workbook(self.output); ws = wb.active
        original = inspect_book(self.source, "Empresas")["rows"]
        for n, row in enumerate(original, 1): self.assertEqual([ws.cell(n,c).value for c in range(1,5)], row)
        self.assertEqual(ws["C2"].value, "Aurora Fictícia")
        self.assertEqual(ws["A2"].value, "0001")
        self.assertEqual(ws["H1"].value, "Pesquisa - Usuários associados")
        self.assertEqual(ws["H2"].value, "123"); self.assertEqual(ws["J2"].value, "51–200 funcionários")
        self.assertEqual(ws["L2"].data_type, "s"); self.assertEqual(ws["L2"].value, "=1+1")
        self.assertIsNone(ws["G6"].value)
        self.assertIsNone(ws["A7"].value)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)
        wb.close()

    def test_cannot_overwrite_original(self):
        for target in (self.source, self.store.job(self.job)["snapshot"]):
            with self.assertRaises(ValueError): write_output(self.store.job(self.job), self.store.records(self.job), target)

    def test_output_lock_keeps_previous_file(self):
        self.output.write_bytes(b"old")
        with patch("pesquisa.excel.os.replace", side_effect=PermissionError("arquivo aberto")):
            with self.assertRaises(PermissionError): write_output(self.store.job(self.job), self.store.records(self.job), self.output)
        self.assertEqual(self.output.read_bytes(), b"old")
        self.assertEqual(list(self.root.glob(".pesquisa-*")), [])

    def test_column_conflict(self):
        jid = self.store.import_job(EXAMPLES / "07_coluna_em_conflito.xlsx", "Empresas", self.output)
        write_output(self.store.job(jid), self.store.records(jid), self.output)
        wb = load_workbook(self.output); headers = [c.value for c in wb.active[1]]
        self.assertIn("Pesquisa - Status (2)", headers)
        self.assertEqual(wb.active["B2"].value, "Texto original deve permanecer"); wb.close()

    def test_persistence_and_optimistic_lock(self):
        d = self.record["data"]; d["name"] = "Salvo"
        self.store.save(self.record["id"], d, 0)
        reopened = Store(self.root / "dados")
        self.assertEqual(reopened.record(self.record["id"])["data"]["name"], "Salvo")
        with self.assertRaises(ValueError): self.store.save(self.record["id"], d, 0)

    def message(self):
        self.store.setting("heartbeat", time.time()); self.store.arm(self.record["id"])
        return dict(action="capture", **self.store.capture_context(), url="https://www.linkedin.com/company/ficticia/about/", confirmed=True, fields={"name":"Fictícia", "members":"1.234", "industry":"Tecnologia", "size":"51–200 funcionários"})

    def test_capture_is_review_draft_and_single_use(self):
        msg = self.message(); response = handle(self.store, msg)
        self.assertTrue(response["ok"])
        self.assertEqual(self.store.record(self.record["id"])["data"]["name"], "")
        self.assertEqual(len(self.store.pending_captures()), 1)
        with self.assertRaises(ValueError): handle(self.store, msg)

    def test_reject_wrong_row_job_expired_revision_unconfirmed(self):
        for key, value in (("record_id", "wrong"), ("job_id", "wrong"), ("token", "wrong"), ("confirmed", False)):
            msg = self.message(); msg[key] = value
            with self.subTest(key=key), self.assertRaises(ValueError): handle(self.store, msg)
        msg = self.message(); self.store.save(self.record["id"], self.record["data"])
        with self.assertRaises(ValueError): handle(self.store, msg)
        msg = self.message()
        with self.store.db() as db: db.execute("UPDATE captures SET expires=0")
        with self.assertRaises(ValueError): handle(self.store, msg)

    def test_heartbeat_required(self):
        msg = self.message(); self.store.setting("heartbeat", 0)
        with self.assertRaises(ValueError): handle(self.store, msg)

    def test_complete_requirements(self):
        d = empty_record("empresa.example"); d["status"] = "Concluída"
        with self.assertRaises(ValueError): validate_record(d)
        d.update(name="Empresa", confirmed=True, url="https://www.linkedin.com/company/test/", unavailable=["members","industry","size"])
        validate_record(d)

    def test_native_framing(self):
        stream = io.BytesIO(); write_message(stream, {"nome":"São João"}); stream.seek(0)
        self.assertEqual(read_message(stream), {"nome":"São João"})
        with self.assertRaises(ValueError): read_message(io.BytesIO(b"\xff\xff\xff\xff"))

class SearchTests(unittest.TestCase):
    def test_request_and_filter(self):
        def opener(request, timeout):
            self.assertEqual(request.full_url, "https://api.tavily.com/search")
            p = json.loads(request.data); self.assertFalse(p["include_raw_content"]); self.assertFalse(p["auto_parameters"])
            return io.BytesIO(json.dumps({"results":[{"url":"https://www.linkedin.com/company/test/about/", "title":"Teste"},{"url":"https://www.linkedin.com/company/test/"},{"url":"https://evil.example/company/test/"}]}).encode())
        self.assertEqual(len(search("FAKE", "query", opener)), 1)

    def test_service_failures(self):
        for code in (401,403,429,432,433,500):
            def opener(*a, **kw): raise HTTPError("https://api.tavily.com/search", code, "error", {}, None)
            with self.subTest(code=code), self.assertRaises(SearchError): search("FAKE", "query", opener)
        def offline(*a, **kw): raise URLError("offline")
        with self.assertRaises(SearchError): search("FAKE", "query", offline)

if __name__ == "__main__": unittest.main()
