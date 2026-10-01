"""Automação com dados fictícios; nenhuma rede ou credencial real."""
import json
import os
import sqlite3
import tempfile
import time
import unittest
import zipfile
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

from pesquisa.store import Store
from pesquisa.backups import create_backup
from pesquisa.ui import App
from pesquisa.search import SearchError

BASE = Path(__file__).resolve().parents[1]
CANDIDATE = dict(url="https://www.linkedin.com/company/ficticia/", title="Fictícia", snippet="Teste")


class AutomationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "dados")
        self.job = self.store.import_job(BASE / "examples/01_empresas_ficticias.xlsx", "Empresas", self.root / "saida.xlsx")
        self.record = self.store.records(self.job)[0]

    def tearDown(self):
        self.temp.cleanup()

    def test_persist_candidates_cache_empty_and_failure(self):
        request = self.store.begin_search(self.record["id"])
        self.assertTrue(self.store.finish_search(request, [CANDIDATE], "2026-10-01T12:00:00"))
        reopened = Store(self.store.root)
        self.assertEqual(reopened.search_info(self.record["id"])["candidates"], [CANDIDATE])
        self.assertEqual(reopened.cached_search(self.job, request[1])[0], [CANDIDATE])
        other = self.store.records(self.job)[1]
        request2 = self.store.begin_search(other["id"])
        self.store.finish_search(request2, [], "today")
        self.assertEqual(reopened.search_info(other["id"])["state"], "Sem candidatas")
        self.store.finish_search(self.store.begin_search(other["id"]), [], "later", "Sem conexão")
        self.assertEqual(reopened.search_info(other["id"])["state"], "Falha na busca")
        self.assertFalse(reopened.record(other["id"])["data"]["confirmed"])

    def test_website_query_and_superseded_request_reject_old_response(self):
        rid = self.record["id"]
        request = self.store.begin_search(rid)
        data = self.record["data"]; data["website"] = "outro.example"
        self.store.save(rid, data)
        self.assertFalse(self.store.finish_search(request, [CANDIDATE], "old"))
        request = self.store.begin_search(rid)
        self.store.set_query(rid, "termos diferentes")
        self.assertFalse(self.store.finish_search(request, [CANDIDATE], "old"))
        request = self.store.begin_search(rid)
        self.store.begin_search(rid)
        self.assertFalse(self.store.finish_search(request, [CANDIDATE], "old"))
        self.assertEqual(self.store.search_info(rid)["candidates"], [])

    def test_migration_backup_and_restore_entries(self):
        with self.store.db() as db:
            db.execute("DROP TABLE searches"); db.execute("DROP TABLE search_cache")
            db.execute("PRAGMA user_version=0")
        migrated = Store(self.store.root)
        self.assertEqual(len(migrated.records(self.job)), 7)
        archive = next((self.store.root / "backups").glob("*-migracao.zip"))
        restored = self.root / "restaurado"
        with zipfile.ZipFile(archive) as source:
            self.assertIsNone(source.testzip()); source.extractall(restored)
        with closing(sqlite3.connect(restored / "trabalhos.sqlite")) as db:
            with db:
                for jid, in db.execute("SELECT id FROM jobs").fetchall():
                    db.execute("UPDATE jobs SET snapshot=?,output=? WHERE id=?", (str(restored / "entradas" / (jid + ".xlsx")), str(restored / "saida.xlsx"), jid))
        store = Store(restored)
        self.assertEqual(store.record(self.record["id"])["original"], self.record["original"])
        self.assertTrue(Path(store.job(self.job)["snapshot"]).exists())

    def test_backup_daily_retention_and_failure(self):
        initial = create_backup(self.store.root)
        self.assertEqual(initial, create_backup(self.store.root))
        for _ in range(8): create_backup(self.store.root, migration=True)
        self.assertEqual(len(list((self.store.root / "backups").glob("*.zip"))), 7)
        with patch("pesquisa.store.create_backup", side_effect=OSError("disco cheio")):
            data = self.store.record(self.record["id"])["data"]; data["notes"] = "Salvo apesar da falha"
            self.store.save(self.record["id"], data)
            self.assertIn("disco cheio", self.store.backup_error)
        self.assertEqual(self.store.record(self.record["id"])["data"]["notes"], "Salvo apesar da falha")

    def test_failed_migration_backup_prevents_schema_changes(self):
        with self.store.db() as db: db.execute("PRAGMA user_version=0")
        with patch("pesquisa.store.create_backup", side_effect=OSError("disco cheio")):
            with self.assertRaises(OSError): Store(self.store.root)
        with self.store.db() as db: self.assertEqual(db.execute("PRAGMA user_version").fetchone()[0], 0)


class AutomationUITests(unittest.TestCase):
    def setUp(self):
        AutomationTests.setUp(self)
        self.messages = patch("pesquisa.ui.messagebox.showerror").start()
        self.info = patch("pesquisa.ui.messagebox.showinfo").start()
        self.app = App(self.store); self.app.withdraw(); self.app.open_job(self.job); self.app.update()

    def tearDown(self):
        self.app.close(); patch.stopall(); AutomationTests.tearDown(self)

    def drain(self):
        deadline = time.monotonic() + 5
        while self.app.searching and time.monotonic() < deadline:
            time.sleep(.01)
            self.app.after_cancel(self.app.poll_job); self.app.poll()
        self.assertFalse(self.app.searching)

    def test_resume_position_and_candidates_without_network(self):
        with patch("pesquisa.ui.read_key", return_value="test"), patch("pesquisa.ui.search", return_value=[CANDIDATE]):
            self.app.search_one(); self.drain()
        rid = self.app.current["id"]
        second = self.store.records(self.job)[1]["id"]
        self.app.load_record(second); self.app.home()
        with patch("pesquisa.ui.search") as search:
            self.app.open_job(self.job)
            self.assertEqual(self.app.current["id"], second)
            self.assertEqual(self.app.results[rid], [CANDIDATE]); search.assert_not_called()

    def test_pending_skips_success_empty_failure_and_finished(self):
        rows = self.store.records(self.job)
        for row, candidates, error in [(rows[0], [CANDIDATE], ""), (rows[1], [], ""), (rows[2], [], "quota")]:
            self.store.finish_search(self.store.begin_search(row["id"]), candidates, "today", error)
        self.app.load_record(rows[0]["id"])
        with patch.object(self.app, "start_search") as start:
            self.app.search_batch()
            self.assertFalse({r["id"] for r in start.call_args.args[0]} & {r["id"] for r in rows[:3]})
            self.app.retry_searches()
            self.assertEqual([r["id"] for r in start.call_args.args[0]], [rows[2]["id"]])

    def test_import_starts_search_after_preview_and_output_choice(self):
        def descendants(widget):
            for child in widget.winfo_children():
                yield child
                yield from descendants(child)
        with patch("pesquisa.ui.filedialog.askopenfilename", return_value=str(BASE / "examples/01_empresas_ficticias.xlsx")), patch("pesquisa.ui.filedialog.asksaveasfilename", return_value=str(self.root / "nova.xlsx")), patch.object(self.app, "search_batch") as search:
            self.app.import_dialog()
            button = next(w for w in descendants(self.app) if w.winfo_class() == "TButton" and w.cget("text") == "Escolher Excel de saída e iniciar")
            search.assert_not_called()
            button.invoke()
            search.assert_called_once()
            self.assertNotEqual(self.app.job_id, self.job)

    @unittest.skipUnless(os.name == "nt", "Bloqueio real exige Windows")
    def test_real_windows_sharing_lock_schedules_retry(self):
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = kernel.CreateFileW(str(self.root / "saida.xlsx"), 0x80000000, 1, None, 3, 0, None)
        self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
        try:
            self.app.vars["notes"].set("teste do bloqueio real")
            self.app.save_next()
            self.assertIn(self.job, self.app.export_retries)
            self.messages.assert_not_called()
        finally:
            kernel.CloseHandle(handle)
        self.app.export_retries[self.job] = (0, str(self.root / "saida.xlsx"))
        self.app.retry_exports()
        self.assertFalse(self.store.job(self.job)["excel_pending"])

    def test_persistent_duplicate_reuse_without_key(self):
        first, second = self.store.records(self.job)[:2]
        data = second["data"]; data["website"] = first["data"]["website"]
        self.store.save(second["id"], data)
        with patch("pesquisa.ui.read_key", return_value="test"), patch("pesquisa.ui.search", return_value=[CANDIDATE]):
            self.app.search_one(); self.drain()
        with patch("pesquisa.ui.read_key") as key, patch("pesquisa.ui.search") as search:
            self.app.start_search([self.store.record(second["id"])]); self.drain()
            key.assert_not_called(); search.assert_not_called()
        self.assertEqual(self.store.search_info(second["id"])["candidates"], [CANDIDATE])
        self.assertFalse(self.store.record(second["id"])["data"]["confirmed"])

    def test_response_after_edit_does_not_fill_row(self):
        import threading
        started = threading.Event(); release = threading.Event()
        def delayed(*args):
            started.set(); release.wait(3); return [CANDIDATE]
        with patch("pesquisa.ui.read_key", return_value="test"), patch("pesquisa.ui.search", side_effect=delayed):
            self.app.search_one(); self.assertTrue(started.wait(2))
            self.app.queryvar.set("outros termos")
            release.set(); self.drain()
        self.assertEqual(self.store.search_info(self.app.current["id"])["candidates"], [])

    def test_locked_excel_advances_and_retry_uses_latest_values(self):
        locked = PermissionError("Excel aberto"); locked.winerror = 32
        rid = self.app.current["id"]
        self.app.vars["notes"].set("primeira alteração")
        with patch("pesquisa.ui.write_output", side_effect=locked): self.app.save_next()
        self.assertNotEqual(self.app.current["id"], rid)
        self.assertTrue(self.store.job(self.job)["excel_pending"])
        self.messages.assert_not_called()
        self.app.vars["notes"].set("alteração mais recente")
        self.app.export_retries[self.job] = (0, self.store.job(self.job)["output"])
        self.app.retry_exports()
        self.assertFalse(self.store.job(self.job)["excel_pending"])
        from openpyxl import load_workbook
        book = load_workbook(self.root / "saida.xlsx")
        self.assertEqual(book.active.cell(self.app.current["row_num"], 12).value, "alteração mais recente")
        book.close()

    def test_next_skips_finished_wraps_and_never_selects_same(self):
        rows = self.store.records(self.job)
        for row in rows[1:]:
            data = row["data"]; data["status"] = "Não encontrada"; self.store.save(row["id"], data)
        self.app.load_record(rows[-1]["id"]); self.app.save_next()
        self.assertEqual(self.app.current["id"], rows[0]["id"])
        self.app.save_next()
        self.assertEqual(self.app.current["id"], rows[0]["id"])
        self.assertIn("única", self.app.progress.get())

    def test_interruption_keeps_received_result_and_unsearched_pending(self):
        def first(*args):
            self.app.stop.set(); return [CANDIDATE]
        with patch("pesquisa.ui.read_key", return_value="test"), patch("pesquisa.ui.search", side_effect=first) as search:
            self.app.search_batch(); self.drain(); self.assertEqual(search.call_count, 1)
        self.assertEqual(self.store.search_info(self.record["id"])["candidates"], [CANDIDATE])
        self.assertTrue(any(self.store.search_info(r["id"])["state"] == "Não executada" for r in self.store.records(self.job)[1:]))
