"""Importa valores e exporta sempre para um arquivo diferente da entrada."""
import os
import tempfile
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .core import FIELDS

HEADER = "Website da empresa"

def inspect_book(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Escolha um arquivo .xlsx.")
    with path.open("rb") as stream:
        wb = load_workbook(stream, data_only=False)
    try:
        if sheet is None:
            return wb.sheetnames
        ws = wb[sheet]
        if ws.max_row > 10000 or ws.max_column > 500:
            raise ValueError("A aba tem uma área excessiva. Copie a tabela para uma aba simples, com até 50 empresas.")
        all_rows = list(ws.iter_rows())
        last_row = max((c.row for row in all_rows for c in row if c.value is not None), default=1)
        last_col = max((c.column for row in all_rows for c in row if c.value is not None), default=1)
        headers = [ws.cell(1, i).value for i in range(1, last_col + 1)]
        matches = [i for i, h in enumerate(headers) if str(h or "").strip() == HEADER]
        if len(matches) != 1:
            raise ValueError('A primeira linha deve ter uma única coluna "Website da empresa".')
        values_wb = load_workbook(path, data_only=True)
        try:
            vals = values_wb[sheet]
            rows, records = [], []
            for n in range(1, last_row + 1):
                row = []
                for c in range(1, last_col + 1):
                    cell = ws.cell(n, c)
                    value = vals.cell(n, c).value
                    if cell.data_type == "f" and value is None:
                        raise ValueError(f"Fórmula sem resultado em {cell.coordinate}. Recalcule e salve no Excel antes de importar.")
                    row.append(value)
                rows.append(row)
                if n > 1 and any(v is not None for v in row):
                    records.append((n, row[matches[0]]))
            if not records:
                raise ValueError("A aba não contém empresas.")
            if len(records) > 50:
                raise ValueError(f"Foram encontradas {len(records)} empresas. O limite é 50; nenhuma linha foi importada.")
            return dict(headers=headers, rows=rows, records=records, sheet=sheet)
        finally:
            values_wb.close()
    finally:
        wb.close()

def same_path(a, b):
    a, b = Path(a), Path(b)
    if os.path.normcase(str(a.resolve())) == os.path.normcase(str(b.resolve())):
        return True
    return a.exists() and b.exists() and os.path.samefile(a, b)

def write_output(job, records, destination, protected_paths=()):
    destination = Path(destination).resolve()
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("O destino deve terminar em .xlsx.")
    for source in (job["source"], job["snapshot"], *protected_paths):
        if same_path(destination, source):
            raise ValueError("O Excel original e as cópias de entrada não podem ser substituídos.")
    source = inspect_book(job["snapshot"], job["sheet"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    for row in source["rows"]:
        ws.append(row)
    extra = ["Website utilizado", "URL LinkedIn confirmada", *FIELDS.values(), "Campos não disponíveis", "Observações", "Status", "Data da consulta"]
    names = {str(v) for v in source["headers"]}
    start = len(source["headers"]) + 1
    for i, label in enumerate(extra, start):
        base = "Pesquisa - " + label
        name, suffix = base, 2
        while name in names:
            name = f"{base} ({suffix})"
            suffix += 1
        names.add(name)
        ws.cell(1, i, name)
    for record in records:
        d = record["data"]
        date = datetime.fromisoformat(d["consulted_at"]).replace(tzinfo=None) if d.get("consulted_at") else None
        values = [d["website"], d["url"] if d.get("confirmed") else "", *(d[k] for k in FIELDS), ", ".join(FIELDS[k] for k in d["unavailable"]), d["notes"], d["status"], date]
        for col, value in enumerate(values, start):
            ws.cell(record["row_num"], col, value)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Arial", size=10)
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24415B")
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 30 if col != start + 1 else 55
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    fd, temp = tempfile.mkstemp(prefix=".pesquisa-", suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    try:
        wb.save(temp)
        os.replace(temp, destination)
    finally:
        wb.close()
        if os.path.exists(temp):
            os.unlink(temp)
