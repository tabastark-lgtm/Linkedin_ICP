"""Gera exclusivamente dados fictícios e casos de erro solicitados."""
from pathlib import Path
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1] / "examples"

def create(name, headers, rows, second=False):
    wb = Workbook(); ws = wb.active; ws.title = "Empresas"
    ws.append(headers)
    for row in rows: ws.append(row)
    if second: wb.create_sheet("Outra aba").append(["Conteúdo fora da importação"])
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24415B")
    for row in list(ws)[1:]:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str) and not cell.value.startswith("="):
                cell.number_format = "@"
    for column, width in [("A",12),("B",45),("C",40),("D",65)]: ws.column_dimensions[column].width = width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
    path = ROOT / name; wb.save(path); wb.close()
    return path

def main():
    ROOT.mkdir(exist_ok=True)
    headers = ["Código", "Website da empresa", "Empresa fictícia", "Caso de teste"]
    create("01_empresas_ficticias.xlsx", headers, [
        ["0001", "www.aurora-exemplo.example", "Aurora Fictícia", "Domínio válido; não se espera página real"],
        ["0002", None, "Estrela Fictícia", "Website ausente"],
        ["0003", "não é uma URL", "Caju Fictícia", "Website inválido"],
        ["0004", "contato@empresa.example", "Horizonte Fictícia", "E-mail não é website"],
        ["0005", "https://aurora-exemplo.example/sobre", "Aurora Fictícia repetida", "Duplicata mantida como registro independente"],
        [None, None, None, None],
        ["0006", "https://oceano-exemplo.example", "Oceano Fictícia", "Acentos: São João; zero à esquerda no código"],
        ["0007", "http://cedro-exemplo.example", "Cedro Fictícia", "Protocolo HTTP aceito para pesquisa"],
    ], second=True)
    create("02_sem_cabecalho.xlsx", ["Código", "Site"], [["001", "empresa.example"]])
    create("03_cabecalho_duplicado.xlsx", ["Website da empresa", "Website da empresa"], [["a.example", "b.example"]])
    for count in (50, 51):
        create(f"04_limite_{count}.xlsx", headers, [[f"{i:04d}", f"empresa-{i}.example", f"Fictícia {i}", f"Limite de {count}"] for i in range(1, count+1)])
    create("05_formula_sem_resultado.xlsx", ["Website da empresa", "Valor calculado"], [["empresa.example", "=1+1"]])
    path = create("06_formula_com_resultado.xlsx", ["Website da empresa", "Valor calculado"], [["empresa.example", "=1+1"]])
    with zipfile.ZipFile(path) as source:
        files = {n: source.read(n) for n in source.namelist()}
    files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dest:
        for name, content in files.items(): dest.writestr(name, content)
    create("07_coluna_em_conflito.xlsx", ["Website da empresa", "Pesquisa - Status"], [["empresa.example", "Texto original deve permanecer"]])
    print("8 exemplos criados em", ROOT)

if __name__ == "__main__": main()
